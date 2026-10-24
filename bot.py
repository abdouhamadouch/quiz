"""Telegram quiz bot with saved quizzes, native quiz polls, and timed exams."""
from __future__ import annotations

import asyncio
import io
import logging
import mimetypes
import os
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

from telegram import InlineKeyboardButton, InlineKeyboardMarkup, Update
from telegram.constants import ChatType
from telegram.ext import (
    Application, CallbackQueryHandler, CommandHandler, ContextTypes,
    MessageHandler, PollAnswerHandler, filters,
)

from ai_provider import extract_text_from_image, generate_questions
from db import (
    copy_quiz, delete_group_draft, delete_quiz as db_delete_quiz, get_group_drafts,
    get_quiz, init_db, leaderboard, list_quizzes, quiz_leaderboard, rename_quiz,
    reset_chat, save_group_draft, save_quiz, save_quiz_results, save_session_scores,
    set_quiz_category, update_quiz_questions, user_stats,
)

logging.basicConfig(level=os.getenv("LOG_LEVEL", "INFO"), format="%(asctime)s %(levelname)s %(name)s %(message)s")
log = logging.getLogger("quizbot")
BOT_TOKEN = os.getenv("BOT_TOKEN", "").strip()
OWNER_IDS = {int(v) for v in os.getenv("OWNER_IDS", "").split(",") if v.strip().isdigit()}
MAX_QUESTIONS = 20
MIN_SECONDS, MAX_SECONDS = 10, 60
sessions: dict[int, dict] = {}
pending_setup: dict[tuple[int, int], dict] = {}
group_drafts: dict[int, dict] = {}
pending_question_edits: dict[tuple[int, int], dict] = {}
sessions_lock = asyncio.Lock()


def group_chat(chat_type: str) -> bool:
    return chat_type in (ChatType.GROUP, ChatType.SUPERGROUP)


async def authorized(update: Update, context: ContextTypes.DEFAULT_TYPE) -> bool:
    user, chat = update.effective_user, update.effective_chat
    if not user or not chat:
        return False
    if user.id in OWNER_IDS:
        return True
    # A user controls their own private chat and private quiz session.
    if chat.type == ChatType.PRIVATE:
        return True
    if not group_chat(chat.type):
        return False
    try:
        member = await context.bot.get_chat_member(chat.id, user.id)
        return member.status in ("administrator", "creator")
    except Exception:
        return False


def config_keyboard(cfg: dict) -> InlineKeyboardMarkup:
    levels = {"easy": "سهل", "medium": "متوسط", "hard": "صعب"}
    modes = {"review_manual": "مراجعة: سؤال بسؤال", "review_all": "نشر كل الاستفتاءات", "exam": "امتحان بالنقاط"}
    rows = [
        [InlineKeyboardButton(f"عدد الأسئلة: {cfg['count']}", callback_data="cfg:count")],
        [InlineKeyboardButton(f"الصعوبة: {levels[cfg['difficulty']]}", callback_data="cfg:difficulty")],
        [InlineKeyboardButton(f"الوقت: {cfg['seconds']} ثانية", callback_data="cfg:time")],
        [InlineKeyboardButton(f"الطريقة: {modes[cfg['mode']]}", callback_data="cfg:mode")],
        [InlineKeyboardButton("إنشاء الاختبار وحفظه", callback_data="cfg:generate")],
        [InlineKeyboardButton("إلغاء", callback_data="cfg:cancel")],
    ]
    return InlineKeyboardMarkup(rows)


def config_text(cfg: dict) -> str:
    topic = cfg.get("topic") or "غير محدد"
    source_state = "محتوى مرفق" if cfg.get("source") else "موضوع عام"
    return (
        "إعداد الاختبار قبل الإنشاء\n"
        f"الموضوع: {topic}\nالمصدر: {source_state}\n"
        "غيّر الخيارات من الأزرار، ثم اضغط «إنشاء الاختبار وحفظه»."
    )


async def reply_lines_in_chunks(message, lines: list[str], limit: int = 3500) -> None:
    chunk = ""
    for line in lines:
        candidate = f"{chunk}\n{line}" if chunk else line
        if len(candidate) > limit and chunk:
            await message.reply_text(chunk)
            chunk = line
        else:
            chunk = candidate
    if chunk:
        await message.reply_text(chunk)


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    msg, chat, user = update.effective_message, update.effective_chat, update.effective_user
    if msg and chat and user and context.args and context.args[0].startswith("share_") and group_chat(chat.type):
        try:
            _, quiz_raw, owner_raw = context.args[0].split("_", 2)
            quiz_id, owner_id = int(quiz_raw), int(owner_raw)
        except (ValueError, TypeError):
            await msg.reply_text("رابط مشاركة الاختبار غير صالح.")
            return
        if owner_id != user.id:
            await msg.reply_text("افتح رابط المشاركة من الحساب الذي أنشأ الاختبار.")
            return
        saved = get_quiz(quiz_id, owner_id)
        if not saved:
            await msg.reply_text("لم أجد هذا الاختبار في مكتبتك.")
            return
        group_drafts[chat.id] = {"quiz": saved, "owner_id": owner_id}
        save_group_draft(chat.id, owner_id, saved)
        await msg.reply_text(
            f"تم تجهيز «{saved['title']}» ({len(saved['questions'])} سؤالًا) كمسودة في هذه المجموعة. "
            "لم يبدأ الامتحان بعد؛ على المالك أو أحد مشرفي المجموعة كتابة /startquiz."
        )
        return
    await update.effective_message.reply_text(
        "🎓 بوت الاختبارات التعليمية\n\n"
        "/quiz — إنشاء اختبار (يمكنك الرد على رسالة درس أو إدخال موضوع ومصدر)\n"
        "/library — اختباراتك المحفوظة\n"
        "/sharequiz ID — نشر اختبار محفوظ في المجموعة كمسودة امتحان\n"
        "/startquiz — بدء الامتحان المنشور في المجموعة\n"
        "/reviewquiz ID — مراجعة باستفتاءات متتابعة\n"
        "/publishquiz ID — نشر جميع الاستفتاءات دفعة واحدة\n"
        "/end أو /stop — إنهاء الاختبار النشط\n"
        "/cancel — إلغاء الاختبار النشط\n"
        "/leaderboard — لوحة المتصدرين\n"
        "/searchquiz كلمة — البحث في المكتبة (في الخاص)\n"
        "/previewquiz ID — معاينة اختبار محفوظ\n"
        "/deletequiz ID — حذف بعد التأكيد\n"
        "/copyquiz ID — نسخ اختبار\n"
        "/renamequiz ID العنوان — تغيير العنوان\n"
        "/editquiz ID رقم_السؤال — تعديل سؤال\n"
        "/quizresults ID — نتائج اختبار المجموعة\n"
        "/setcategory ID تصنيف — تصنيف الاختبار\n"
        "/stats — إحصائياتك الشخصية"
    )


async def extract_source_message(message, context: ContextTypes.DEFAULT_TYPE) -> str:
    """Extract text from Telegram text, PDF/DOCX/TXT files, or image OCR."""
    direct = (message.text or "").strip()
    caption = (message.caption or "").strip()
    if direct:
        return direct[:14000]
    if caption and not (message.photo or message.document):
        return caption[:14000]
    file_id, filename, mime_type = None, "", ""
    if message.photo:
        file_id = message.photo[-1].file_id
        filename, mime_type = "photo.jpg", "image/jpeg"
    elif message.document:
        file_id = message.document.file_id
        filename = message.document.file_name or "upload"
        mime_type = message.document.mime_type or mimetypes.guess_type(filename)[0] or ""
    else:
        raise ValueError("لم أجد نصًا أو ملفًا مدعومًا في الرسالة.")
    if message.document and message.document.file_size and message.document.file_size > 15 * 1024 * 1024:
        raise ValueError("حجم الملف أكبر من 15 ميغابايت.")
    telegram_file = await context.bot.get_file(file_id)
    data = bytes(await telegram_file.download_as_bytearray())
    suffix = filename.lower()
    if mime_type.startswith("image/") or suffix.endswith((".jpg", ".jpeg", ".png", ".webp")):
        image_mime = mime_type if mime_type.startswith("image/") else mimetypes.guess_type(filename)[0] or "image/jpeg"
        extracted_image = await asyncio.to_thread(extract_text_from_image, data, image_mime)
        return (caption + "\n\n" + extracted_image).strip()[:14000] if caption else extracted_image
    if suffix.endswith(".pdf") or mime_type == "application/pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if len(reader.pages) > 40:
            raise ValueError("ملفات PDF محدودة بـ 40 صفحة في المرة الواحدة.")
        extracted = "\n".join(page.extract_text() or "" for page in reader.pages).strip()
        if not extracted:
            raise ValueError("لم أجد نصًا قابلًا للاستخراج في PDF؛ إذا كان ممسوحًا ضوئيًا أرسل الصفحات كصور.")
        return (caption + "\n\n" + extracted).strip()[:14000] if caption else extracted[:14000]
    if suffix.endswith(".docx") or "wordprocessingml" in mime_type:
        from docx import Document
        document = Document(io.BytesIO(data))
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
        extracted = "\n".join(paragraphs).strip()
        if not extracted:
            raise ValueError("ملف DOCX لا يحتوي على نص قابل للقراءة.")
        return (caption + "\n\n" + extracted).strip()[:14000] if caption else extracted[:14000]
    if suffix.endswith(".txt") or mime_type.startswith("text/"):
        extracted_text = data.decode("utf-8-sig", errors="replace").strip()
        return (caption + "\n\n" + extracted_text).strip()[:14000] if caption else extracted_text[:14000]
    raise ValueError("نوع الملف غير مدعوم. استخدم PDF أو DOCX أو TXT أو صورة.")


async def collect_setup_media(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    msg, chat, user = update.effective_message, update.effective_chat, update.effective_user
    if not msg or not chat or not user:
        return
    cfg = pending_setup.get((chat.id, user.id))
    if not cfg:
        return
    if cfg.get("stage") != "source":
        await msg.reply_text("أرسل الموضوع أولًا، ثم أرفق ملف الدرس أو صورته في خطوة المصدر.")
        return
    try:
        cfg["source"] = await extract_source_message(msg, context)
    except ValueError as exc:
        await msg.reply_text(f"تعذّر قراءة الملف: {str(exc)[:250]}")
        return
    except Exception:
        log.error("Source file extraction failed: %s", type(exc).__name__)
        await msg.reply_text("تعذّر قراءة الملف. تحقق من نوعه وحجمه، ومن إعداد مفتاح Gemini إذا كان المصدر صورة.")
        return
    cfg.pop("stage", None)
    sent = await msg.reply_text(config_text(cfg), reply_markup=config_keyboard(cfg))
    cfg["config_message_id"] = sent.message_id


async def quiz(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    msg, chat, user = update.effective_message, update.effective_chat, update.effective_user
    if not msg or not chat or not user:
        return
    if group_chat(chat.type) and not await authorized(update, context):
        await msg.reply_text("إنشاء الاختبارات الجماعية متاح للمالك أو مشرفي المجموعة فقط. يمكنك إنشاء الاختبار في الخاص ثم مشاركته.")
        return
    source = ""
    if msg.reply_to_message:
        try:
            source = await extract_source_message(msg.reply_to_message, context)
        except ValueError as exc:
            await msg.reply_text(f"تعذّر قراءة المصدر: {str(exc)[:250]}\nأرسل نصًا أو ملف PDF/DOCX أو صورة واضحة.")
            return
        except Exception:
            log.error("Replied source extraction failed: %s", type(exc).__name__)
            await msg.reply_text("تعذّر قراءة المصدر. تحقق من الملف ومفتاح Gemini إذا كان المصدر صورة.")
            return
    cfg = {
        "chat_id": chat.id, "chat_type": chat.type, "creator_id": user.id,
        "source": source.strip(), "topic": "", "count": 5, "seconds": 20,
        "difficulty": "medium", "mode": "exam",
    }
    pending_setup[(chat.id, user.id)] = cfg
    if source:
        sent = await msg.reply_text(config_text(cfg), reply_markup=config_keyboard(cfg))
        cfg["config_message_id"] = sent.message_id
    else:
        cfg["stage"] = "topic"
        await msg.reply_text("أرسل موضوع الاختبار، أو /skip لتجاوز الموضوع.")
        

async def collect_setup_text(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    msg, chat, user = update.effective_message, update.effective_chat, update.effective_user
    if not msg or not chat or not user:
        return
    edit_key = (chat.id, user.id)
    edit = pending_question_edits.get(edit_key)
    if edit:
        lines = (msg.text or "").splitlines()
        if len(lines) < 6:
            await msg.reply_text("التنسيق غير مكتمل. أرسل نص السؤال، ثم 4 خيارات، ثم رقم الإجابة الصحيحة من 1 إلى 4، ثم التفسير اختياريًا.")
            return
        question = lines[0].strip()
        options = [line.strip() for line in lines[1:5]]
        try:
            answer_index = int(lines[5].strip()) - 1
        except ValueError:
            await msg.reply_text("رقم الإجابة يجب أن يكون من 1 إلى 4.")
            return
        if not question or any(not option for option in options) or len({option.casefold() for option in options}) != 4 or answer_index not in range(4):
            await msg.reply_text("السؤال أو الخيارات غير صالحة. تأكد من وجود 4 خيارات مختلفة ورقم إجابة من 1 إلى 4.")
            return
        saved = get_quiz(edit["quiz_id"], user.id)
        if not saved:
            pending_question_edits.pop(edit_key, None)
            await msg.reply_text("لم أعد أجد هذا الاختبار في مكتبتك.")
            return
        explanation = "\n".join(lines[6:]).strip()[:700] if len(lines) > 6 else ""
        saved["questions"][edit["index"]] = {
            "question": question[:900], "options": [option[:160] for option in options],
            "answer": answer_index, "explanation": explanation,
        }
        if update_quiz_questions(edit["quiz_id"], user.id, saved["questions"]):
            pending_question_edits.pop(edit_key, None)
            await msg.reply_text(f"تم تحديث السؤال {edit['index'] + 1} في الاختبار #{edit['quiz_id']}. استخدم /previewquiz {edit['quiz_id']} للمراجعة.")
        else:
            await msg.reply_text("تعذّر حفظ التعديل. حاول مرة أخرى.")
        return
    key = (chat.id, user.id)
    cfg = pending_setup.get(key)
    if not cfg:
        return
    value = (msg.text or "").strip()
    if cfg.get("stage") == "topic":
        cfg["topic"] = value[:300]
        cfg["stage"] = "source"
        await msg.reply_text("أرسل نص الدرس أو الكلمات التي تريد الاعتماد عليها، أو /skip لإنشاء أسئلة من الموضوع فقط.")
        return
    if cfg.get("stage") == "source":
        cfg["source"] = value[:14000]
        cfg.pop("stage", None)
        sent = await msg.reply_text(config_text(cfg), reply_markup=config_keyboard(cfg))
        cfg["config_message_id"] = sent.message_id


async def skip_setup(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    chat, user, msg = update.effective_chat, update.effective_user, update.effective_message
    if not chat or not user or not msg:
        return
    cfg = pending_setup.get((chat.id, user.id))
    if not cfg:
        await msg.reply_text("لا توجد خطوة إعداد تنتظر التخطي.")
        return
    if cfg.get("stage") == "topic":
        cfg["topic"] = ""
        cfg["stage"] = "source"
        await msg.reply_text("أرسل نص الدرس أو الكلمات، أو /skip لإنشاء أسئلة من المعرفة العامة حول الموضوع.")
    elif cfg.get("stage") == "source":
        cfg["source"] = ""
        cfg.pop("stage", None)
        sent = await msg.reply_text(config_text(cfg), reply_markup=config_keyboard(cfg))
        cfg["config_message_id"] = sent.message_id
    else:
        await msg.reply_text("الإعداد جاهز. استخدم الأزرار لاختيار الخيارات.")


async def configure(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    query = update.callback_query
    if not query:
        return
    parts = query.data.split(":")
    if len(parts) < 2:
        await query.answer("طلب غير صالح.", show_alert=True)
        return
    action = parts[1]
    key = next((k for k, v in pending_setup.items()
                if v.get("config_message_id") == query.message.message_id and k[0] == query.message.chat_id), None)
    if not key:
        await query.answer("انتهت جلسة الإعداد.", show_alert=True)
        await query.edit_message_text("انتهت جلسة الإعداد. أرسل /quiz للبدء من جديد.")
        return
    cfg = pending_setup[key]
    if query.from_user.id != cfg.get("creator_id"):
        await query.answer("هذه الإعدادات تخص الشخص الذي بدأها.", show_alert=True)
        return
    if group_chat(query.message.chat.type) and action == "generate" and not await authorized(update, context):
        await query.answer("إنشاء الاختبار الجماعي متاح للمالك أو مشرفي المجموعة فقط.", show_alert=True)
        return
    await query.answer()
    if action == "cancel":
        pending_setup.pop(key, None)
        await query.edit_message_text("تم إلغاء إعداد الاختبار.")
        return
    if action == "count":
        cfg["count"] = {5: 10, 10: 15, 15: 20, 20: 5}.get(cfg["count"], 5)
    elif action == "difficulty":
        cfg["difficulty"] = {"easy": "medium", "medium": "hard", "hard": "easy"}[cfg["difficulty"]]
    elif action == "time":
        cfg["seconds"] = {10: 20, 20: 30, 30: 45, 45: 60, 60: 10}.get(cfg["seconds"], 20)
    elif action == "mode":
        cfg["mode"] = {"exam": "review_manual", "review_manual": "review_all", "review_all": "exam"}[cfg["mode"]]
    elif action == "generate":
        pending_setup.pop(key, None)
        await query.edit_message_text("جارٍ إنشاء الأسئلة والتحقق منها...")
        source = cfg["source"]
        if not source:
            source = f"General educational topic: {cfg['topic'] or 'English language and grammar'}. Create accurate general-knowledge questions on this topic."
        elif cfg["topic"]:
            source = f"Focus topic: {cfg['topic']}\n\nSource material:\n{source}"
        try:
            questions, provider = await asyncio.to_thread(
                generate_questions, source, cfg["count"], cfg["topic"], cfg["difficulty"]
            )
        except Exception:
            log.exception("AI generation failed")
            await query.message.reply_text("تعذّر إنشاء الأسئلة. تحقق من مفاتيح Groq وGemini ثم حاول مجددًا.")
            return
        title = cfg["topic"] or (source[:70].replace("\n", " ") if cfg["source"] else "اختبار عام")
        quiz_id = save_quiz(cfg["creator_id"], title, source, questions, cfg["seconds"], cfg["mode"], cfg["difficulty"], cfg["topic"] or "عام")
        if cfg["chat_type"] == ChatType.PRIVATE:
            bot = await context.bot.get_me()
            share_url = f"https://t.me/{bot.username}?startgroup=share_{quiz_id}_{cfg['creator_id']}"
            choice_keyboard = InlineKeyboardMarkup([
                [InlineKeyboardButton("ابدأ الاختبار في الخاص", callback_data=f"startprivate:{quiz_id}:{cfg['seconds']}:{cfg['mode']}")],
                [InlineKeyboardButton("انشر الاختبار في مجموعة", url=share_url)],
            ])
            await query.message.reply_text(
                f"حُفظ الاختبار في مكتبتك برقم #{quiz_id}، وعدد أسئلته {len(questions)}. المزود: {provider}.\n\nأين تريد استخدامه؟",
                reply_markup=choice_keyboard,
            )
        else:
            await query.message.reply_text(
                f"حُفظ الاختبار في مكتبتك برقم #{quiz_id}، وعدد أسئلته {len(questions)}. المزود: {provider}."
            )
            await launch_quiz(context, cfg["chat_id"], cfg["chat_type"], cfg["creator_id"],
                              questions, cfg["seconds"], cfg["mode"], quiz_id)
        return
    await query.edit_message_text(config_text(cfg), reply_markup=config_keyboard(cfg))


async def start_saved_private(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    query = update.callback_query
    if not query or query.message is None or query.message.chat.type != ChatType.PRIVATE:
        return
    try:
        _, quiz_raw, seconds_raw, mode = query.data.split(":", 3)
        quiz_id, seconds = int(quiz_raw), int(seconds_raw)
    except (ValueError, AttributeError):
        await query.answer("اختيار غير صالح.", show_alert=True)
        return
    saved = get_quiz(quiz_id, query.from_user.id)
    if not saved:
        await query.answer("لم أجد الاختبار في مكتبتك.", show_alert=True)
        return
    if query.message.chat_id in sessions:
        await query.answer("يوجد اختبار نشط بالفعل في الخاص.", show_alert=True)
        return
    await query.answer("بدء الاختبار")
    await query.edit_message_reply_markup(reply_markup=None)
    await launch_quiz(context, query.message.chat_id, ChatType.PRIVATE, query.from_user.id,
                      saved["questions"], seconds, mode, quiz_id)


async def launch_quiz(context: ContextTypes.DEFAULT_TYPE, chat_id: int, chat_type: str,
                      creator_id: int, questions: list[dict], seconds: int,
                      mode: str, quiz_id: int | None = None) -> None:
    if chat_id in sessions:
        await context.bot.send_message(chat_id, "يوجد اختبار نشط في هذه المحادثة. أنهِه أولًا باستخدام /end.")
        return
    session = {
        "chat_id": chat_id, "chat_type": chat_type, "creator_id": creator_id,
        "questions": questions, "seconds": seconds, "index": 0, "scores": {},
        "answered": set(), "message_id": None, "timer": None, "mode": mode,
        "poll_ids": [], "quiz_id": quiz_id,
    }
    sessions[chat_id] = session
    if mode == "review_all":
        await publish_all_polls(context, session)
    elif mode == "review_manual":
        await post_review_poll(context, session)
    else:
        await context.bot.send_message(chat_id, "بدأ الامتحان التنافسي. أجب عن الأسئلة قبل انتهاء الوقت.")
        await post_exam_question(context, session)


async def post_review_poll(context: ContextTypes.DEFAULT_TYPE, session: dict) -> None:
    index = session["index"]
    questions = session["questions"]
    if index >= len(questions):
        sessions.pop(session["chat_id"], None)
        await context.bot.send_message(session["chat_id"], "انتهت المراجعة بالاستفتاءات.")
        return
    q = questions[index]
    poll = await context.bot.send_poll(
        chat_id=session["chat_id"],
        question=f"{index + 1}/{len(questions)}. {q['question']}"[:300],
        options=[str(o)[:100] for o in q["options"]],
        type="quiz", correct_option_id=q["answer"], is_anonymous=False,
        explanation=q.get("explanation", "")[:200],
    )
    session["poll_ids"].append(poll.message_id)
    keyboard = InlineKeyboardMarkup([[InlineKeyboardButton("السؤال التالي ⏭", callback_data=f"pollnext:{session['chat_id']}:{index}")],
                                     [InlineKeyboardButton("إنهاء المراجعة", callback_data=f"pollend:{session['chat_id']}:{index}")]])
    await context.bot.send_message(session["chat_id"], f"الاستفتاء {index + 1}/{len(questions)}. انتقل يدويًا عند الاستعداد.", reply_markup=keyboard)


async def publish_all_polls(context: ContextTypes.DEFAULT_TYPE, session: dict) -> None:
    chat_id = session["chat_id"]
    for index, q in enumerate(session["questions"]):
        try:
            poll = await context.bot.send_poll(
                chat_id=chat_id, question=f"{index + 1}/{len(session['questions'])}. {q['question']}"[:300],
                options=[str(o)[:100] for o in q["options"]], type="quiz",
                correct_option_id=q["answer"], is_anonymous=False,
                explanation=q.get("explanation", "")[:200],
            )
            session["poll_ids"].append(poll.message_id)
        except Exception:
            log.exception("Could not publish poll %s", index + 1)
    sessions.pop(chat_id, None)
    await context.bot.send_message(chat_id, f"تم نشر {len(session['poll_ids'])} استفتاءات دفعة واحدة. تظهر الإجابة الصحيحة بعد تصويت كل مشارك.")


async def poll_navigation(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    query = update.callback_query
    if not query:
        return
    try:
        action, chat_raw, index_raw = query.data.split(":")
        chat_id, index = int(chat_raw), int(index_raw)
    except (ValueError, AttributeError):
        await query.answer("زر غير صالح.", show_alert=True)
        return
    session = sessions.get(chat_id)
    if not session or session.get("mode") != "review_manual" or session["index"] != index:
        await query.answer("انتهت هذه الخطوة.", show_alert=True)
        return
    if query.from_user.id != session["creator_id"] and not await authorized(update, context):
        await query.answer("الانتقال متاح لمنشئ المراجعة أو مشرف المجموعة.", show_alert=True)
        return
    await query.answer()
    if action == "pollend":
        sessions.pop(chat_id, None)
        await query.edit_message_text("تم إنهاء المراجعة.")
        return
    if session["poll_ids"]:
        try:
            await context.bot.stop_poll(chat_id, session["poll_ids"][-1])
        except Exception:
            pass
    session["index"] += 1
    await query.edit_message_reply_markup(reply_markup=None)
    await post_review_poll(context, session)


async def poll_answer(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    # Telegram quiz polls reveal correctness directly to each voter. This handler is
    # retained for future per-participant review analytics; no group-wide vote count is assumed.
    answer = update.poll_answer
    if answer:
        log.debug("Poll answer received for poll %s", answer.poll_id)


async def post_exam_question(context: ContextTypes.DEFAULT_TYPE, session: dict) -> None:
    chat_id, index = session["chat_id"], session["index"]
    if index >= len(session["questions"]):
        await finish(context, chat_id, cancelled=False)
        return
    q = session["questions"][index]
    session["answered"] = set()
    keyboard = [[InlineKeyboardButton(option, callback_data=f"ans:{chat_id}:{index}:{i}")]
                for i, option in enumerate(q["options"])]
    try:
        sent = await context.bot.send_message(
            chat_id, f"السؤال {index + 1}/{len(session['questions'])}\nالوقت: {session['seconds']} ثانية\n\n{q['question']}",
            reply_markup=InlineKeyboardMarkup(keyboard),
        )
        session["message_id"] = sent.message_id
    except Exception:
        log.exception("Unable to post exam question")
        sessions.pop(chat_id, None)
        return
    session["timer"] = asyncio.create_task(timeout_question(context, chat_id, index, session["seconds"]))


async def timeout_question(context: ContextTypes.DEFAULT_TYPE, chat_id: int, index: int, seconds: int) -> None:
    try:
        await asyncio.sleep(seconds)
        session = sessions.get(chat_id)
        if not session or session["index"] != index or session.get("mode") != "exam":
            return
        q = session["questions"][index]
        try:
            await context.bot.edit_message_reply_markup(chat_id, session["message_id"], reply_markup=None)
            await context.bot.send_message(chat_id, f"انتهى الوقت. الإجابة الصحيحة: {q['options'][q['answer']]}\n{q.get('explanation', '')}".strip())
        except Exception:
            log.exception("Unable to send timeout feedback")
        session["index"] += 1
        await asyncio.sleep(1)
        await post_exam_question(context, session)
    except asyncio.CancelledError:
        return


async def answer(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    query = update.callback_query
    if not query:
        return
    try:
        _, chat_raw, index_raw, option_raw = query.data.split(":")
        chat_id, index, option = int(chat_raw), int(index_raw), int(option_raw)
    except (ValueError, AttributeError):
        await query.answer("زر غير صالح.", show_alert=True)
        return
    session = sessions.get(chat_id)
    if not session or session.get("mode") != "exam" or session["index"] != index:
        await query.answer("انتهى هذا السؤال.", show_alert=True)
        return
    user = query.from_user
    if user.id in session["answered"]:
        await query.answer("سُجّلت إجابتك بالفعل.")
        return
    q = session["questions"][index]
    if option not in range(4):
        await query.answer("اختيار غير صالح.", show_alert=True)
        return
    session["answered"].add(user.id)
    score = session["scores"].setdefault(user.id, {"name": user.full_name, "points": 0, "correct": 0, "answered": 0})
    score["name"], score["answered"] = user.full_name, score["answered"] + 1
    correct = option == q["answer"]
    if correct:
        score["points"] += 10
        score["correct"] += 1
    await query.answer("صحيح! +10 نقاط" if correct else f"خطأ. الصحيح: {q['options'][q['answer']]}", show_alert=True)
    if group_chat(session["chat_type"]):
        try:
            await context.bot.send_message(chat_id, f"{user.full_name}: {'إجابة صحيحة (+10)' if correct else 'إجابة غير صحيحة'}")
        except Exception:
            log.exception("Unable to post group feedback")
    if session["chat_type"] == ChatType.PRIVATE:
        timer = session.get("timer")
        if timer:
            timer.cancel()
        try:
            await context.bot.edit_message_reply_markup(chat_id, session["message_id"], reply_markup=None)
            await context.bot.send_message(chat_id, f"{'إجابة صحيحة' if correct else 'الإجابة الصحيحة: ' + q['options'][q['answer']]}\n{q.get('explanation', '')}".strip())
        except Exception:
            log.exception("Unable to send private feedback")
        session["index"] += 1
        await post_exam_question(context, session)


async def finish(context: ContextTypes.DEFAULT_TYPE, chat_id: int, cancelled: bool) -> None:
    session = sessions.pop(chat_id, None)
    if not session:
        return
    timer = session.get("timer")
    if timer and timer is not asyncio.current_task():
        timer.cancel()
    if session.get("mode", "exam") != "exam":
        await context.bot.send_message(chat_id, "أُنهيت جلسة الاستفتاءات." if cancelled else "انتهت جلسة الاستفتاءات.")
        return
    if cancelled:
        await context.bot.send_message(chat_id, "أُلغي الامتحان ولم تُحفظ نقاط هذه الجلسة.")
        return
    save_session_scores(chat_id, session["scores"])
    if session.get("quiz_id") is not None:
        save_quiz_results(chat_id, int(session["quiz_id"]), session["scores"])
    participants = sorted(session["scores"].values(), key=lambda x: (x["points"], x["correct"]), reverse=True)
    lines = ["النتائج النهائية", f"عدد الأسئلة: {len(session['questions'])}"]
    if not participants:
        lines.append("لم تُسجّل أي إجابات.")
    else:
        for i, row in enumerate(participants[:10], 1):
            lines.append(f"{i}. {row['name']} — {row['points']} نقطة ({row['correct']}/{row['answered']} صحيحة)")
    await context.bot.send_message(chat_id, "\n".join(lines))


async def end_quiz(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not await authorized(update, context):
        await update.effective_message.reply_text("هذا الأمر متاح للمالك أو مشرفي المجموعة فقط.")
        return
    if update.effective_chat.id not in sessions:
        await update.effective_message.reply_text("لا يوجد اختبار نشط.")
        return
    await finish(context, update.effective_chat.id, cancelled=False)


async def cancel_quiz(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not await authorized(update, context):
        await update.effective_message.reply_text("هذا الأمر متاح للمالك أو مشرفي المجموعة فقط.")
        return
    chat_id = update.effective_chat.id
    if chat_id not in sessions:
        if chat_id in group_drafts:
            group_drafts.pop(chat_id, None)
            delete_group_draft(chat_id)
            await update.effective_message.reply_text("تم إلغاء مسودة الامتحان المنشورة.")
        else:
            pending_setup.pop((chat_id, update.effective_user.id), None)
            pending_question_edits.pop((chat_id, update.effective_user.id), None)
            await update.effective_message.reply_text("أُلغي أي إعداد أو تعديل شخصي قيد الانتظار.")
        return
    await finish(context, chat_id, cancelled=True)


async def search_quizzes(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    user, msg, chat = update.effective_user, update.effective_message, update.effective_chat
    if not user or not msg or not chat:
        return
    if chat.type != ChatType.PRIVATE:
        await msg.reply_text("استخدم /searchquiz في الخاص لحماية خصوصية مكتبتك.")
        return
    term = " ".join(context.args).strip()
    if not term:
        await msg.reply_text("الصيغة: /searchquiz كلمة أو جزء من العنوان")
        return
    rows = list_quizzes(user.id, limit=30, search=term)
    if not rows:
        await msg.reply_text("لم أعثر على اختبارات تطابق البحث.")
        return
    await reply_lines_in_chunks(msg, ["نتائج البحث:"] + [
        f"#{row['id']} — {row['title']} [{row.get('category', 'عام')}] ({row['question_count']} سؤالًا)" for row in rows
    ])


async def preview_quiz(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    user, msg, chat = update.effective_user, update.effective_message, update.effective_chat
    if not user or not msg or not chat:
        return
    if chat.type != ChatType.PRIVATE:
        await msg.reply_text("المعاينة متاحة في الخاص فقط.")
        return
    if not context.args or not context.args[0].isdigit():
        await msg.reply_text("الصيغة: /previewquiz ID")
        return
    saved = get_quiz(int(context.args[0]), user.id)
    if not saved:
        await msg.reply_text("لم أجد هذا الاختبار في مكتبتك.")
        return
    blocks = [f"معاينة الاختبار #{saved['id']}: {saved['title']}",
              f"الأسئلة: {len(saved['questions'])} | الصعوبة: {saved.get('difficulty', 'medium')} | الوقت: {saved.get('seconds', 20)} ثانية"]
    for index, question in enumerate(saved["questions"], 1):
        blocks.append(
            f"\n{index}. {question['question']}\n" +
            "\n".join(f"{n}. {option}" for n, option in enumerate(question["options"], 1)) +
            f"\nالإجابة: {question['answer'] + 1}. {question['options'][question['answer']]}\nالتفسير: {question.get('explanation', '')}"
        )
    chunks, current = [], ""
    for block in blocks:
        if len(current) + len(block) > 3500:
            chunks.append(current)
            current = ""
        current += block + "\n"
    if current:
        chunks.append(current)
    for chunk in chunks:
        await msg.reply_text(chunk)
    await msg.reply_text(
        "إدارة هذا الاختبار:\n/editquiz ID رقم_السؤال — تعديل سؤال\n"
        "/renamequiz ID العنوان — تغيير الاسم\n/copyquiz ID — إنشاء نسخة\n/deletequiz ID — حذف بعد التأكيد"
    )


async def delete_quiz_command(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    user, msg, chat = update.effective_user, update.effective_message, update.effective_chat
    if not user or not msg or not chat:
        return
    if chat.type != ChatType.PRIVATE:
        await msg.reply_text("حذف الاختبارات متاح في الخاص فقط.")
        return
    if not context.args or not context.args[0].isdigit():
        await msg.reply_text("الصيغة: /deletequiz ID")
        return
    quiz_id = int(context.args[0])
    saved = get_quiz(quiz_id, user.id)
    if not saved:
        await msg.reply_text("الاختبار غير موجود في مكتبتك أو لا تملك صلاحية حذفه.")
        return
    keyboard = InlineKeyboardMarkup([[
        InlineKeyboardButton("تأكيد الحذف", callback_data=f"quizdelete:confirm:{quiz_id}"),
        InlineKeyboardButton("إلغاء", callback_data=f"quizdelete:cancel:{quiz_id}"),
    ]])
    await msg.reply_text(f"هل تريد حذف الاختبار #{quiz_id} «{saved['title']}» نهائيًا؟", reply_markup=keyboard)


async def confirm_delete_quiz(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    query = update.callback_query
    if not query:
        return
    try:
        _, action, raw_id = query.data.split(":")
        quiz_id = int(raw_id)
    except (ValueError, AttributeError):
        await query.answer("طلب غير صالح.", show_alert=True)
        return
    if query.message is None or query.message.chat.type != ChatType.PRIVATE:
        await query.answer("افتح هذا الأمر في الخاص.", show_alert=True)
        return
    if action == "cancel":
        await query.answer("تم الإلغاء.")
        await query.edit_message_text("أُلغي حذف الاختبار.")
        return
    if not get_quiz(quiz_id, query.from_user.id):
        await query.answer("لا تملك هذا الاختبار أو أنه حُذف بالفعل.", show_alert=True)
        return
    if db_delete_quiz(quiz_id, query.from_user.id):
        await query.answer("تم الحذف.")
        await query.edit_message_text(f"تم حذف الاختبار #{quiz_id} من مكتبتك.")
    else:
        await query.answer("تعذّر حذف الاختبار.", show_alert=True)


async def copy_saved_quiz(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    user, msg, chat = update.effective_user, update.effective_message, update.effective_chat
    if not user or not msg or not chat:
        return
    if chat.type != ChatType.PRIVATE:
        await msg.reply_text("نسخ الاختبارات متاح في الخاص فقط.")
        return
    if not context.args or not context.args[0].isdigit():
        await msg.reply_text("الصيغة: /copyquiz ID")
        return
    new_id = copy_quiz(int(context.args[0]), user.id)
    if new_id is None:
        await msg.reply_text("لم أجد الاختبار في مكتبتك.")
    else:
        await msg.reply_text(f"تم إنشاء نسخة مستقلة برقم #{new_id}.")


async def rename_saved_quiz(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    user, msg, chat = update.effective_user, update.effective_message, update.effective_chat
    if not user or not msg or not chat:
        return
    if chat.type != ChatType.PRIVATE:
        await msg.reply_text("تغيير العنوان متاح في الخاص فقط.")
        return
    if len(context.args) < 2 or not context.args[0].isdigit():
        await msg.reply_text("الصيغة: /renamequiz ID العنوان الجديد")
        return
    quiz_id, title = int(context.args[0]), " ".join(context.args[1:]).strip()
    if not title:
        await msg.reply_text("اكتب عنوانًا جديدًا.")
        return
    if rename_quiz(quiz_id, user.id, title):
        await msg.reply_text(f"تم تغيير عنوان الاختبار #{quiz_id}.")
    else:
        await msg.reply_text("لم أجد الاختبار في مكتبتك.")


async def edit_saved_question(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    user, msg, chat = update.effective_user, update.effective_message, update.effective_chat
    if not user or not msg or not chat:
        return
    if chat.type != ChatType.PRIVATE:
        await msg.reply_text("تعديل الأسئلة متاح في الخاص فقط.")
        return
    if len(context.args) != 2 or not all(part.isdigit() for part in context.args):
        await msg.reply_text("الصيغة: /editquiz ID رقم_السؤال")
        return
    quiz_id, number = map(int, context.args)
    saved = get_quiz(quiz_id, user.id)
    if not saved or not 1 <= number <= len(saved["questions"]):
        await msg.reply_text("لم أجد الاختبار أو رقم السؤال خارج النطاق.")
        return
    pending_question_edits[(chat.id, user.id)] = {"quiz_id": quiz_id, "index": number - 1}
    q = saved["questions"][number - 1]
    await msg.reply_text(
        f"تعديل السؤال {number} في الاختبار #{quiz_id}.\n"
        "أرسل 6 أسطر بهذا الترتيب، ويمكن إضافة التفسير كسطر سابع:\n"
        "1) نص السؤال\n2) الخيار الأول\n3) الخيار الثاني\n4) الخيار الثالث\n5) الخيار الرابع\n"
        "6) رقم الإجابة الصحيحة من 1 إلى 4\n7) التفسير (اختياري، ويمكن تركه فارغًا)\n\n"
        f"السؤال الحالي: {q['question']}\nأرسل /cancel لإلغاء التعديل."
    )


async def show_quiz_results(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    msg, chat = update.effective_message, update.effective_chat
    if not msg or not chat:
        return
    if not group_chat(chat.type):
        await msg.reply_text("استخدم /quizresults ID داخل المجموعة التي أُجري فيها الاختبار.")
        return
    if not context.args or not context.args[0].isdigit():
        await msg.reply_text("الصيغة: /quizresults ID")
        return
    rows = quiz_leaderboard(chat.id, int(context.args[0]))
    if not rows:
        await msg.reply_text("لا توجد نتائج محفوظة لهذا الاختبار في هذه المجموعة.")
        return
    await msg.reply_text("نتائج الاختبار\n" + "\n".join(
        f"{index}. {row['name']} — {row['points']} نقطة ({row['correct']}/{row['answered']} صحيحة)"
        for index, row in enumerate(rows, 1)
    ))


async def set_quiz_category_command(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    user, msg, chat = update.effective_user, update.effective_message, update.effective_chat
    if not user or not msg or not chat:
        return
    if chat.type != ChatType.PRIVATE:
        await msg.reply_text("تصنيف المكتبة متاح في الخاص فقط.")
        return
    if len(context.args) < 2 or not context.args[0].isdigit():
        await msg.reply_text("الصيغة: /setcategory ID التصنيف")
        return
    quiz_id, category = int(context.args[0]), " ".join(context.args[1:])
    if set_quiz_category(quiz_id, user.id, category):
        await msg.reply_text(f"تم تعيين تصنيف «{category[:80]}» للاختبار #{quiz_id}.")
    else:
        await msg.reply_text("لم أجد الاختبار في مكتبتك.")


async def show_stats(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    user, msg, chat = update.effective_user, update.effective_message, update.effective_chat
    if not user or not msg or not chat:
        return
    if chat.type != ChatType.PRIVATE:
        await msg.reply_text("افتح البوت في الخاص لعرض إحصائياتك.")
        return
    stats = user_stats(user.id)
    accuracy = round(100 * stats["correct"] / stats["answered"]) if stats["answered"] else 0
    await msg.reply_text(
        "إحصائياتك الشخصية\n"
        f"الاختبارات المحفوظة: {stats['quiz_count']}\n"
        f"المجموعات التي شاركت فيها: {stats['groups_count']}\n"
        f"مجموع النقاط: {stats['points']}\n"
        f"الإجابات الصحيحة: {stats['correct']} من {stats['answered']}\n"
        f"نسبة الإجابات الصحيحة: {accuracy}%"
    )


async def show_library(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    user = update.effective_user
    if not user:
        return
    if update.effective_chat.type != ChatType.PRIVATE:
        await update.effective_message.reply_text("افتح البوت في الخاص لإدارة مكتبتك.")
        return
    quizzes = list_quizzes(user.id)
    if not quizzes:
        await update.effective_message.reply_text("مكتبتك فارغة. أنشئ اختبارًا باستخدام /quiz.")
        return
    lines = ["اختباراتك المحفوظة (الأحدث أولًا):"]
    lines.extend(f"#{q['id']} — {q['title']} [{q.get('category', 'عام')}] ({q['question_count']} سؤالًا)" for q in quizzes)
    lines.append("/previewquiz ID — معاينة | /searchquiz كلمة — بحث")
    lines.append("/deletequiz ID — حذف | /copyquiz ID — نسخ")
    lines.append("/renamequiz ID العنوان — تغيير الاسم")
    lines.append("/editquiz ID رقم_السؤال — تعديل سؤال")
    lines.append("/setcategory ID تصنيف — تصنيف")
    lines.append("/sharequiz ID — مشاركة | /reviewquiz ID — مراجعة")
    lines.append("/stats — إحصائياتي")
    await reply_lines_in_chunks(update.effective_message, lines)


async def share_quiz(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    msg, chat, user = update.effective_message, update.effective_chat, update.effective_user
    if not msg or not chat or not user or not group_chat(chat.type):
        if msg:
            await msg.reply_text("استخدم /sharequiz ID داخل المجموعة التي تريد نشر الاختبار فيها.")
        return
    if not context.args or not context.args[0].isdigit():
        await msg.reply_text("الصيغة: /sharequiz ID")
        return
    saved = get_quiz(int(context.args[0]), user.id)
    if not saved:
        await msg.reply_text("لم أجد اختبارًا بهذا الرقم في مكتبتك.")
        return
    group_drafts[chat.id] = {"quiz": saved, "owner_id": user.id}
    save_group_draft(chat.id, user.id, saved)
    await msg.reply_text(
        f"تم تجهيز «{saved['title']}» ({len(saved['questions'])} سؤالًا) كمسودة امتحان. "
        "لم يبدأ بعد. اكتب /startquiz لبدئه أو /cancel لإلغائه."
    )


async def start_shared_quiz(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    msg, chat, user = update.effective_message, update.effective_chat, update.effective_user
    if not msg or not chat or not user:
        return
    if not await authorized(update, context):
        await msg.reply_text("بدء الامتحان في المجموعة متاح للمالك أو مشرفي المجموعة.")
        return
    draft = group_drafts.get(chat.id)
    if not draft:
        await msg.reply_text("لا توجد مسودة منشورة. استخدم /sharequiz ID أولًا.")
        return
    if chat.id in sessions:
        await msg.reply_text("يوجد اختبار نشط بالفعل.")
        return
    group_drafts.pop(chat.id, None)
    delete_group_draft(chat.id)
    saved = draft["quiz"]
    await launch_quiz(context, chat.id, chat.type, draft["owner_id"], saved["questions"], int(saved.get("seconds", 20)), "exam", saved["id"])


async def launch_saved_review(update: Update, context: ContextTypes.DEFAULT_TYPE, mode: str) -> None:
    msg, chat, user = update.effective_message, update.effective_chat, update.effective_user
    if not msg or not chat or not user:
        return
    if group_chat(chat.type) and not await authorized(update, context):
        await msg.reply_text("نشر الاختبارات في المجموعة متاح للمالك أو مشرفي المجموعة فقط.")
        return
    if not context.args or not context.args[0].isdigit():
        await msg.reply_text(f"الصيغة: /{'reviewquiz' if mode == 'review_manual' else 'publishquiz'} ID")
        return
    saved = get_quiz(int(context.args[0]), user.id)
    if not saved:
        await msg.reply_text("لم أجد اختبارًا بهذا الرقم في مكتبتك.")
        return
    if chat.id in sessions:
        await msg.reply_text("يوجد اختبار نشط في هذه المحادثة. أنهِه أولًا.")
        return
    await launch_quiz(context, chat.id, chat.type, user.id, saved["questions"], int(saved.get("seconds", 20)), mode, saved["id"])


async def review_quiz(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    await launch_saved_review(update, context, "review_manual")


async def publish_quiz(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    await launch_saved_review(update, context, "review_all")


async def reset(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not await authorized(update, context):
        await update.effective_message.reply_text("هذا الأمر متاح للمالك أو مشرفي المجموعة فقط.")
        return
    chat_id = update.effective_chat.id
    if chat_id in sessions:
        await finish(context, chat_id, cancelled=True)
    reset_chat(chat_id)
    await update.effective_message.reply_text("تم مسح نتائج هذه المحادثة وإعادة ضبطها.")


async def show_leaderboard(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    rows = leaderboard(update.effective_chat.id)
    if not rows:
        await update.effective_message.reply_text("لا توجد نتائج محفوظة بعد.")
        return
    await update.effective_message.reply_text(
        "لوحة المتصدرين\n" + "\n".join(
            f"{i}. {r['name']} — {r['points']} نقطة | {r['correct']}/{r['answered']} صحيحة"
            for i, r in enumerate(rows, 1)
        )
    )


def start_health_server() -> None:
    port = int(os.getenv("PORT", "10000"))
    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            if self.path not in ("/", "/health"):
                self.send_response(404)
                self.end_headers()
                return
            body = b'{"status":"ok","service":"telegram-quiz-bot"}'
            self.send_response(200)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
        def log_message(self, *_args):
            return
    server = ThreadingHTTPServer(("0.0.0.0", port), Handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()


def main() -> None:
    if not BOT_TOKEN:
        raise RuntimeError("BOT_TOKEN environment variable is required")
    init_db()
    group_drafts.update(get_group_drafts())
    start_health_server()
    app = Application.builder().token(BOT_TOKEN).build()
    app.add_handler(CommandHandler(["start", "help"], start))
    app.add_handler(CommandHandler("quiz", quiz))
    app.add_handler(CommandHandler("skip", skip_setup))
    app.add_handler(CommandHandler("library", show_library))
    app.add_handler(CommandHandler("searchquiz", search_quizzes))
    app.add_handler(CommandHandler("previewquiz", preview_quiz))
    app.add_handler(CommandHandler("deletequiz", delete_quiz_command))
    app.add_handler(CommandHandler("copyquiz", copy_saved_quiz))
    app.add_handler(CommandHandler("renamequiz", rename_saved_quiz))
    app.add_handler(CommandHandler("editquiz", edit_saved_question))
    app.add_handler(CommandHandler("quizresults", show_quiz_results))
    app.add_handler(CommandHandler("setcategory", set_quiz_category_command))
    app.add_handler(CommandHandler("stats", show_stats))
    app.add_handler(CommandHandler("sharequiz", share_quiz))
    app.add_handler(CommandHandler("startquiz", start_shared_quiz))
    app.add_handler(CommandHandler("reviewquiz", review_quiz))
    app.add_handler(CommandHandler("publishquiz", publish_quiz))
    app.add_handler(CommandHandler(["end", "stop"], end_quiz))
    app.add_handler(CommandHandler("cancel", cancel_quiz))
    app.add_handler(CommandHandler("reset", reset))
    app.add_handler(CommandHandler("leaderboard", show_leaderboard))
    app.add_handler(CallbackQueryHandler(configure, pattern=r"^cfg:"))
    app.add_handler(CallbackQueryHandler(start_saved_private, pattern=r"^startprivate:"))
    app.add_handler(CallbackQueryHandler(confirm_delete_quiz, pattern=r"^quizdelete:(confirm|cancel):"))
    app.add_handler(CallbackQueryHandler(poll_navigation, pattern=r"^poll(next|end):"))
    app.add_handler(CallbackQueryHandler(answer, pattern=r"^ans:"))
    app.add_handler(PollAnswerHandler(poll_answer))
    app.add_handler(MessageHandler(filters.Regex(r"^(اختبار|أنشئ اختبار|انشئ اختبار)$"), quiz))
    app.add_handler(MessageHandler(filters.Regex(r"^(مكتبتي|اختباراتي)$"), show_library))
    app.add_handler(MessageHandler(filters.Regex(r"^(المتصدرون|لوحة المتصدرين)$"), show_leaderboard))
    app.add_handler(MessageHandler(filters.Regex(r"^(إحصائياتي|احصائياتي)$"), show_stats))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, collect_setup_text))
    app.add_handler(MessageHandler((filters.Document.ALL | filters.PHOTO) & ~filters.COMMAND, collect_setup_media))
    log.info("Starting Telegram quiz bot polling")
    app.run_polling(drop_pending_updates=False)


if __name__ == "__main__":
    main()
